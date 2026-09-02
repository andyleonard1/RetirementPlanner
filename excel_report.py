"""
Excel Report

Creates the Retirement Planner workbook.
"""

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


class ExcelReport:

    def generate(self, timeline, assumptions):

        workbook = Workbook()

        #
        # Worksheets
        #
        summary = workbook.active
        summary.title = "Summary"

        assumptions_ws = workbook.create_sheet("Assumptions")
        cashflow_ws = workbook.create_sheet("Cash Flow")
        workbook.create_sheet("Tax")
        workbook.create_sheet("Assets")
        workbook.create_sheet("Charts")

        #
        # Fonts
        #
        header_font = Font(bold=True)

        #
        # Summary Sheet
        #
        summary["A1"] = "Retirement Planner"
        summary["A1"].font = Font(size=16, bold=True)

        summary["A2"] = "Version 0.5.0"

        summary["A4"] = "Years Modelled"
        summary["B4"] = len(timeline)

        summary["A5"] = "Starting Pension"
        summary["B5"] = assumptions.get("starting_pension")

        summary["A6"] = "Final Pension"
        summary["B6"] = timeline[-1].closing_pension

        summary["A7"] = "Final ISA"
        summary["B7"] = timeline[-1].isa_closing

        summary["A8"] = "Final Savings"
        summary["B8"] = timeline[-1].savings_closing

        summary["A9"] = "Total Assets"
        summary["B9"] = timeline[-1].total_assets

        #
        # Assumptions Sheet
        #
        assumptions_ws["A1"] = "Assumption"
        assumptions_ws["B1"] = "Value"

        assumptions_ws["A1"].font = header_font
        assumptions_ws["B1"].font = header_font

        row = 2

        for key, value in assumptions.data.items():

            assumptions_ws.cell(row=row, column=1).value = key
            assumptions_ws.cell(row=row, column=2).value = value

            row += 1

        assumptions_ws.freeze_panes = "A2"
        assumptions_ws.auto_filter.ref = assumptions_ws.dimensions

        #
        # Cash Flow Sheet
        #
        headers = [
            "Age",
            "Year",
            "Opening Pension",
            "Growth",
            "Gross Withdrawal",
            "Income Tax",
            "Net Withdrawal",
            "Your State Pension",
            "Spouse State Pension",
            "Cash Available",
            "Savings",
            "ISA",
            "Total Assets",
        ]

        for col, header in enumerate(headers, start=1):
            cell = cashflow_ws.cell(row=1, column=col)
            cell.value = header
            cell.font = header_font

        row = 2

        for year in timeline:

            cashflow_ws.cell(row=row, column=1).value = year.age
            cashflow_ws.cell(row=row, column=2).value = year.calendar_year
            cashflow_ws.cell(row=row, column=3).value = year.opening_pension
            cashflow_ws.cell(row=row, column=4).value = year.pension_growth
            cashflow_ws.cell(row=row, column=5).value = year.gross_pension_income
            cashflow_ws.cell(row=row, column=6).value = year.income_tax
            cashflow_ws.cell(row=row, column=7).value = year.net_pension_income
            cashflow_ws.cell(row=row, column=8).value = year.your_state_pension
            cashflow_ws.cell(row=row, column=9).value = year.spouse_state_pension
            cashflow_ws.cell(row=row, column=10).value = year.cash_available
            cashflow_ws.cell(row=row, column=11).value = year.savings_closing
            cashflow_ws.cell(row=row, column=12).value = year.isa_closing
            cashflow_ws.cell(row=row, column=13).value = year.total_assets

            row += 1

        cashflow_ws.freeze_panes = "A2"
        cashflow_ws.auto_filter.ref = cashflow_ws.dimensions

        #
        # Currency formatting
        #
        currency = "£#,##0.00"

        for ws in [summary, cashflow_ws]:

            for row in ws.iter_rows():

                for cell in row:

                    if isinstance(cell.value, (int, float)):
                        cell.number_format = currency

        #
        # Age and Year should remain integers
        #
        for row in range(2, cashflow_ws.max_row + 1):
            cashflow_ws.cell(row=row, column=1).number_format = "0"
            cashflow_ws.cell(row=row, column=2).number_format = "0"

        #
        # Auto-size every worksheet
        #
        for ws in workbook.worksheets:

            for column_cells in ws.columns:

                length = max(
                    len(str(cell.value if cell.value is not None else ""))
                    for cell in column_cells
                )

                ws.column_dimensions[
                    get_column_letter(column_cells[0].column)
                ].width = length + 3

        #
        # Save workbook
        #
        workbook.save("RetirementPlanner.xlsx")

        print()
        print("Excel report written to RetirementPlanner.xlsx")