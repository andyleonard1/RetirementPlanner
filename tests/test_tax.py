"""
Excel Report

Creates the Retirement Planner workbook.
"""

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference

class ExcelReport:

    def generate(self, timeline, assumptions):

        self.timeline = timeline
        self.assumptions = assumptions

        self.workbook = Workbook()

        self.summary_ws = self.workbook.active
        self.summary_ws.title = "Summary"

        self.assumptions_ws = self.workbook.create_sheet("Assumptions")
        self.cashflow_ws = self.workbook.create_sheet("Cash Flow")
        self.tax_ws = self.workbook.create_sheet("Tax")
        self.assets_ws = self.workbook.create_sheet("Assets")
        self.charts_ws = self.workbook.create_sheet("Charts")

        self.header_font = Font(bold=True)
        self.currency = "£#,##0.00"

        self.build_summary()
        self.build_assumptions()
        self.build_cashflow()
        self.build_tax()
        self.build_assets()

        self.format_workbook()

        self.workbook.save("RetirementPlanner.xlsx")

        print()
        print("Excel report written to RetirementPlanner.xlsx")

    # ----------------------------------------------------

    def build_summary(self):

        ws = self.summary_ws

        ws["A1"] = "Retirement Planner"
        ws["A1"].font = Font(size=16, bold=True)

        ws["A2"] = "Version 0.5.1"

        ws["A4"] = "Years Modelled"
        ws["B4"] = len(self.timeline)

        ws["A5"] = "Starting Pension"
        ws["B5"] = self.assumptions.get("starting_pension")

        ws["A6"] = "Ending Pension"
        ws["B6"] = self.timeline[-1].closing_pension

        ws["A7"] = "Ending ISA"
        ws["B7"] = self.timeline[-1].isa_closing

        ws["A8"] = "Ending Savings"
        ws["B8"] = self.timeline[-1].savings_closing

        ws["A9"] = "Total Assets"
        ws["B9"] = self.timeline[-1].total_assets

        ws["A10"] = "Total Tax Paid"
        ws["B10"] = sum(y.income_tax for y in self.timeline)

    # ----------------------------------------------------

    def build_assumptions(self):

        ws = self.assumptions_ws

        ws.append(["Assumption", "Value"])

        for cell in ws[1]:
            cell.font = self.header_font

        for key, value in self.assumptions.data.items():
            ws.append([key, value])

    # ----------------------------------------------------

    def build_cashflow(self):

        ws = self.cashflow_ws

        headers = [
            "Age",
            "Year",
            "Opening Pension",
            "Growth",
            "Gross Withdrawal",
            "Income Tax",
            "Net Withdrawal",
            "Your State Pension",
            "Spouse State Pension",
            "Cash Available",
            "Savings",
            "ISA",
            "Total Assets",
        ]

        ws.append(headers)

        for cell in ws[1]:
            cell.font = self.header_font

        for year in self.timeline:

            ws.append([
                year.age,
                year.calendar_year,
                year.opening_pension,
                year.pension_growth,
                year.gross_pension_income,
                year.income_tax,
                year.net_pension_income,
                year.your_state_pension,
                year.spouse_state_pension,
                year.cash_available,
                year.savings_closing,
                year.isa_closing,
                year.total_assets,
            ])

    # ----------------------------------------------------

    def build_tax(self):

        ws = self.tax_ws

        headers = [
            "Age",
            "Year",
            "Gross Pension",
            "Taxable Income",
            "Income Tax",
            "Net Pension",
        ]

        ws.append(headers)

        for cell in ws[1]:
            cell.font = self.header_font

        for year in self.timeline:

            ws.append([
                year.age,
                year.calendar_year,
                year.gross_pension_income,
                year.taxable_pension_income,
                year.income_tax,
                year.net_pension_income,
            ])

        total_row = ws.max_row + 2

        ws.cell(total_row, 4).value = "Total Tax"

        ws.cell(total_row, 5).value = sum(
            y.income_tax for y in self.timeline
        )

    # ----------------------------------------------------

    def build_assets(self):

        ws = self.assets_ws

        headers = [
            "Age",
            "Pension",
            "Savings",
            "ISA",
            "Total Assets",
        ]

        ws.append(headers)

        for cell in ws[1]:
            cell.font = self.header_font

        for year in self.timeline:

            ws.append([
                year.age,
                year.closing_pension,
                year.savings_closing,
                year.isa_closing,
                year.total_assets,
            ])

    # ----------------------------------------------------
def build_charts(self):

    chart = LineChart()
    chart.title = "Pension Value"
    chart.style = 2
    chart.y_axis.title = "£"
    chart.x_axis.title = "Age"

    data = Reference(
        self.assets_ws,
        min_col=2,
        min_row=1,
        max_row=self.assets_ws.max_row,
    )

    categories = Reference(
        self.assets_ws,
        min_col=1,
        min_row=2,
        max_row=self.assets_ws.max_row,
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    self.charts_ws.add_chart(chart, "A1")
        #--------------------------------------------------------
    def format_workbook(self):

        for ws in self.workbook.worksheets:

            ws.freeze_panes = "A2"

            ws.auto_filter.ref = ws.dimensions

            for row in ws.iter_rows():

                for cell in row:

                    if isinstance(cell.value, (int, float)):

                        cell.number_format = self.currency

            if ws.max_row > 1:

                for row in range(2, ws.max_row + 1):

                    ws.cell(row=row, column=1).number_format = "0"

                    if ws.max_column >= 2:
                        ws.cell(row=row, column=2).number_format = "0"

            for column in ws.columns:

                width = max(
                    len(str(cell.value or ""))
                    for cell in column
                )

                ws.column_dimensions[
                    get_column_letter(column[0].column)
                ].width = width + 3